"""Read-only profiling for hospital FAQ workbooks."""

from __future__ import annotations

import re
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import BinaryIO

from openpyxl import load_workbook

from medical_agent.modules.knowledge.schemas import FaqEntryInput

REQUIRED_HEADERS = ("分类", "标准问题", "答案内容")
OPTIONAL_HEADERS = ("其他问法", "有效日期")


@dataclass(frozen=True)
class DuplicateQuestion:
    question: str
    rows: list[int]


@dataclass(frozen=True)
class FaqProfile:
    sheet_name: str
    source_rows: int
    valid_rows: int
    headers: list[str]
    missing_required_headers: list[str]
    blank_question_rows: list[int]
    blank_answer_rows: list[int]
    duplicate_questions: list[DuplicateQuestion]
    category_counts: dict[str, int]
    alias_rows: int
    related_knowledge_cells: int
    invalid_effective_date_rows: list[int]
    warnings: list[str]

    def to_dict(self) -> dict[str, object]:
        return asdict(self)


def _normalize_text(value: object) -> str:
    return " ".join(str(value or "").split()).strip()


def _question_key(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "")).casefold()


def profile_faq_xlsx(source: str | Path | BinaryIO) -> FaqProfile:
    workbook = load_workbook(source, read_only=True, data_only=False)
    if not workbook.sheetnames:
        raise ValueError("workbook_has_no_sheets")

    sheet = workbook[workbook.sheetnames[0]]
    row_iterator = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(row_iterator)
    except StopIteration as exc:
        raise ValueError("worksheet_is_empty") from exc

    headers = [_normalize_text(value) for value in raw_headers]
    index = {header: position for position, header in enumerate(headers) if header}
    missing_headers = [header for header in REQUIRED_HEADERS if header not in index]
    if missing_headers:
        return FaqProfile(
            sheet_name=sheet.title,
            source_rows=0,
            valid_rows=0,
            headers=headers,
            missing_required_headers=missing_headers,
            blank_question_rows=[],
            blank_answer_rows=[],
            duplicate_questions=[],
            category_counts={},
            alias_rows=0,
            related_knowledge_cells=0,
            invalid_effective_date_rows=[],
            warnings=["缺少必需列，未继续解析数据行。"],
        )

    question_rows: dict[str, list[int]] = defaultdict(list)
    question_display: dict[str, str] = {}
    category_counts: Counter[str] = Counter()
    blank_question_rows: list[int] = []
    blank_answer_rows: list[int] = []
    invalid_effective_date_rows: list[int] = []
    alias_rows = 0
    related_knowledge_cells = 0
    source_rows = 0
    valid_rows = 0
    related_indexes = [
        position for position, header in enumerate(headers) if header.startswith("关联知识")
    ]

    for row_number, row in enumerate(row_iterator, start=2):
        if not any(value not in (None, "") for value in row):
            continue
        source_rows += 1
        question = _normalize_text(row[index["标准问题"]] if index["标准问题"] < len(row) else None)
        answer = _normalize_text(row[index["答案内容"]] if index["答案内容"] < len(row) else None)
        category = _normalize_text(row[index["分类"]] if index["分类"] < len(row) else None)

        if not question:
            blank_question_rows.append(row_number)
        if not answer:
            blank_answer_rows.append(row_number)
        if question and answer and category:
            valid_rows += 1

        if question:
            key = _question_key(question)
            question_rows[key].append(row_number)
            question_display.setdefault(key, question)
        if category:
            category_counts[category] += 1

        alias_index = index.get("其他问法")
        if alias_index is not None and alias_index < len(row) and _normalize_text(row[alias_index]):
            alias_rows += 1

        effective_index = index.get("有效日期")
        if effective_index is not None and effective_index < len(row):
            value = row[effective_index]
            if value not in (None, "") and not isinstance(value, (date, datetime)):
                invalid_effective_date_rows.append(row_number)

        related_knowledge_cells += sum(
            1
            for position in related_indexes
            if position < len(row) and _normalize_text(row[position])
        )

    duplicates = [
        DuplicateQuestion(question=question_display[key], rows=rows)
        for key, rows in question_rows.items()
        if len(rows) > 1
    ]
    warnings: list[str] = []
    if len(workbook.sheetnames) > 1:
        warnings.append("当前只检查第一个工作表，其他工作表需要显式选择后再导入。")
    if duplicates:
        warnings.append("存在重复标准问题，发布知识库前需要人工确认保留或合并。")
    if invalid_effective_date_rows:
        warnings.append("有效日期包含非 Excel 日期值，已保留原值且未自动转换。")

    return FaqProfile(
        sheet_name=sheet.title,
        source_rows=source_rows,
        valid_rows=valid_rows,
        headers=headers,
        missing_required_headers=missing_headers,
        blank_question_rows=blank_question_rows,
        blank_answer_rows=blank_answer_rows,
        duplicate_questions=duplicates,
        category_counts=dict(category_counts),
        alias_rows=alias_rows,
        related_knowledge_cells=related_knowledge_cells,
        invalid_effective_date_rows=invalid_effective_date_rows,
        warnings=warnings,
    )


def parse_faq_xlsx(source: str | Path | BinaryIO) -> tuple[list[FaqEntryInput], FaqProfile]:
    """Parse valid FAQ rows without mutating the hospital workbook."""
    profile = profile_faq_xlsx(source)
    if profile.missing_required_headers:
        raise ValueError("faq_required_headers_missing")
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = [_normalize_text(value) for value in next(rows)]
    index = {header: position for position, header in enumerate(headers) if header}
    entries: list[FaqEntryInput] = []
    for row_number, row in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in row):
            continue
        category = _normalize_text(row[index["分类"]])
        question = _normalize_text(row[index["标准问题"]])
        answer = _normalize_text(row[index["答案内容"]])
        if not category or not question or not answer:
            continue
        aliases = ""
        alias_index = index.get("其他问法")
        if alias_index is not None and alias_index < len(row):
            aliases = _normalize_text(row[alias_index])
        effective_date = ""
        effective_index = index.get("有效日期")
        if effective_index is not None and effective_index < len(row):
            effective_date = _normalize_text(row[effective_index])
        entries.append(
            FaqEntryInput(
                category=category,
                question=question,
                aliases=aliases,
                answer=answer,
                effective_date=effective_date,
                source_sheet=sheet.title,
                source_row=row_number,
            )
        )
    return entries, profile
